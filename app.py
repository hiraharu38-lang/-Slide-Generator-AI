import streamlit as st
from google import genai
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pypdf import PdfReader
from docx import Document
import json
import time
import io
import os

st.set_page_config(page_title="Slide Generator AI PRO", page_icon="🎨", layout="wide")

# ============================================================
# 初期設定
# ============================================================
SELECT_MODEL = "gemini-2.5-flash"
api_key = st.secrets["GEMINI_API_KEY"]
client = genai.Client(api_key=api_key)
IMAGE_DIR = "./images/"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# ============================================================
# 配色パレット（プレゼンのテーマに応じてAIが1つ選ぶ）
# ============================================================
PALETTES = {
    "midnight_executive": {"primary": "1E2761", "secondary": "CADCFC", "accent": "FFFFFF", "text_on_light": "11182B"},
    "forest_moss":        {"primary": "2C5F2D", "secondary": "97BC62", "accent": "F5F5F5", "text_on_light": "16321A"},
    "coral_energy":       {"primary": "2F3C7E", "secondary": "F9E795", "accent": "F96167", "text_on_light": "1B1F3B"},
    "ocean_gradient":     {"primary": "21295C", "secondary": "1C7293", "accent": "9FD8DF", "text_on_light": "13193A"},
    "charcoal_minimal":   {"primary": "212121", "secondary": "36454F", "accent": "F2F2F2", "text_on_light": "1A1A1A"},
    "teal_trust":         {"primary": "028090", "secondary": "00A896", "accent": "02C39A", "text_on_light": "07343A"},
    "berry_cream":        {"primary": "6D2E46", "secondary": "A26769", "accent": "ECE2D0", "text_on_light": "3C1827"},
    "cherry_bold":        {"primary": "2F3C7E", "secondary": "990011", "accent": "FCF6F5", "text_on_light": "1B1F3B"},
}

def hex_to_rgb(hex_str):
    return RGBColor.from_string(hex_str)

def get_palette(key):
    return PALETTES.get(key, PALETTES["midnight_executive"])

# ============================================================
# リトライ付きGemini呼び出し
# ============================================================
def generate_with_retry(prompt, max_retries=6):
    for i in range(max_retries):
        try:
            response = client.models.generate_content(
                model=SELECT_MODEL,
                contents=prompt,
                config={"response_mime_type": "application/json"},
            )
            return response.text
        except Exception as e:
            if i < max_retries - 1:
                time.sleep(2.0)
                continue
            raise e

# ============================================================
# ファイルからテキスト抽出（PDF / Word）
# ============================================================
def extract_text(uploaded_file):
    if uploaded_file.type == "application/pdf":
        reader = PdfReader(uploaded_file)
        return "".join([(page.extract_text() or "") for page in reader.pages])
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(uploaded_file)
        return "\n".join([para.text for para in doc.paragraphs])
    return ""

# ============================================================
# 共通ヘルパー
# ============================================================
def set_background(slide, rgb):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb

def add_rect(slide, x, y, w, h, rgb, line=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb
    if line:
        shp.line.color.rgb = rgb
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp

def add_circle(slide, x, y, d, rgb):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp

def add_textbox(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT,
                 italic=False, font="Calibri", anchor=None, line_spacing=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.italic = italic
    p.font.name = font
    p.font.color.rgb = color
    return box

def add_bullets(slide, x, y, w, h, bullets, size, color, accent_rgb, font="Calibri", gap=10):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = font
        p.space_after = Pt(gap)
        p.level = 0
    return box

def add_picture_safe(slide, image_name, box_x, box_y, box_w, box_h):
    """指定したbox内に収まるよう縦横比を保ったまま画像を配置する（はみ出し・重なり防止）"""
    if not image_name or image_name == "none":
        return False
    path = os.path.join(IMAGE_DIR, f"{image_name}.png")
    if not os.path.exists(path):
        return False
    try:
        from PIL import Image
        with Image.open(path) as im:
            iw, ih = im.size
    except Exception:
        iw, ih = 1, 1

    img_aspect = iw / ih
    box_aspect = box_w / box_h

    if img_aspect > box_aspect:
        w = box_w
        h = int(w / img_aspect)
    else:
        h = box_h
        w = int(h * img_aspect)

    x = box_x + int((box_w - w) / 2)
    y = box_y + int((box_h - h) / 2)
    slide.shapes.add_picture(path, x, y, width=w, height=h)
    return True


def fit_title_size(text, base_size, soft_limit, min_size, step=4):
    """長い文字列の場合にフォントサイズを段階的に縮める"""
    if not text:
        return base_size
    size = base_size
    length = len(text)
    while length > soft_limit and size > min_size:
        size -= step
        soft_limit += 6
    return size


def add_circle_with_number(slide, x, y, d, number, circle_rgb, number_rgb, font_size=24):
    """円と数字を完全に同じ矩形・中央揃えで配置（位置ズレ防止）"""
    add_circle(slide, x, y, d, circle_rgb)
    box = slide.shapes.add_textbox(x, y, d, d)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = str(number)
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(font_size)
    p.font.bold = True
    p.font.color.rgb = number_rgb
    p.font.name = "Calibri"
    return box

def add_page_number(slide, n, total, rgb):
    add_textbox(
        slide, SLIDE_W - Inches(0.8), SLIDE_H - Inches(0.55), Inches(0.5), Inches(0.4),
        f"{n}", 12, rgb, align=PP_ALIGN.RIGHT, font="Calibri"
    )

# ============================================================
# レイアウト：表紙
# ============================================================
def layout_title(prs, data, palette, idx, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    primary = hex_to_rgb(palette["primary"])
    accent = hex_to_rgb(palette["accent"])
    secondary = hex_to_rgb(palette["secondary"])

    set_background(slide, primary)

    # 装飾円（モチーフ：右上に大きな円、左下に小さな円）
    add_circle(slide, SLIDE_W - Inches(3.5), Inches(-2.0), Inches(5.5), secondary)
    add_circle(slide, Inches(-1.5), SLIDE_H - Inches(1.8), Inches(3.2), secondary)

    title_text = data.get("title", "Untitled")
    title_size = fit_title_size(title_text, base_size=44, soft_limit=14, min_size=28)
    add_textbox(
        slide, Inches(1.0), Inches(2.5), Inches(11.3), Inches(2.2),
        title_text, title_size, accent, bold=True,
        font="Cambria", line_spacing=1.1
    )
    subtitle = data.get("bullets", [])
    if subtitle:
        add_textbox(
            slide, Inches(1.0), Inches(4.9), Inches(10.0), Inches(0.8),
            subtitle[0], 18, accent, italic=True, font="Calibri"
        )
    add_page_number(slide, idx, total, accent)
    return slide

# ============================================================
# レイアウト：章見出し
# ============================================================
def layout_chapter(prs, data, palette, idx, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    primary = hex_to_rgb(palette["primary"])
    accent = hex_to_rgb(palette["accent"])
    secondary = hex_to_rgb(palette["secondary"])

    set_background(slide, primary)
    add_circle(slide, Inches(9.8), Inches(4.2), Inches(4.5), secondary)

    chapter_no = data.get("chapter_no", "")
    if chapter_no:
        add_textbox(slide, Inches(1.0), Inches(2.3), Inches(3.0), Inches(1.0),
                    f"{chapter_no}", 64, secondary, bold=True, font="Cambria")
    title_text = data.get("title", "")
    title_size = fit_title_size(title_text, base_size=38, soft_limit=14, min_size=26)
    add_textbox(
        slide, Inches(1.0), Inches(3.4), Inches(10.8), Inches(2.0),
        title_text, title_size, accent, bold=True, font="Cambria", line_spacing=1.15
    )
    add_page_number(slide, idx, total, accent)
    return slide

# ============================================================
# レイアウト：本文（テキスト＋画像 半々）
# ============================================================
def layout_content_split(prs, data, palette, idx, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    primary = hex_to_rgb(palette["primary"])
    secondary = hex_to_rgb(palette["secondary"])
    accent = hex_to_rgb(palette["accent"])
    text_dark = hex_to_rgb(palette["text_on_light"])

    set_background(slide, hex_to_rgb("FFFFFF"))

    # 右側のカラーパネル（画像が無くても色面として機能する）
    add_rect(slide, Inches(8.6), 0, SLIDE_W - Inches(8.6), SLIDE_H, primary)

    title_text = data.get("title", "")
    title_size = fit_title_size(title_text, base_size=32, soft_limit=16, min_size=22)
    add_textbox(
        slide, Inches(0.7), Inches(0.55), Inches(7.4), Inches(1.0),
        title_text, title_size, text_dark, bold=True, font="Cambria"
    )
    bullets = data.get("bullets", []) or []
    n = max(len(bullets), 1)
    area_top, area_bottom = Inches(1.8), Inches(7.0)
    slot_h = (area_bottom - area_top) / n
    bullet_color = hex_to_rgb(palette["text_on_light"])
    for i, b in enumerate(bullets):
        y = area_top + slot_h * i
        add_circle(slide, Inches(0.7), y + Inches(0.06), Inches(0.16), primary)
        add_textbox(slide, Inches(1.1), y, Inches(6.9), slot_h - Inches(0.1), b, 18, bullet_color,
                    font="Calibri", line_spacing=1.2, anchor=MSO_ANCHOR.TOP)

    # 画像は右パネル内の専用エリアに限定して配置（テキストとは重ならない）
    img_box = (Inches(9.0), Inches(2.0), Inches(3.4), Inches(4.6))
    has_img = add_picture_safe(slide, data.get("image_name"), *img_box)
    if not has_img:
        add_circle(slide, Inches(10.2), Inches(3.5), Inches(1.7), secondary)

    add_page_number(slide, idx, total, text_dark)
    return slide

# ============================================================
# レイアウト：アイコン行（番号＋見出し＋説明 のリスト形式）
# ============================================================
def layout_content_icons(prs, data, palette, idx, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    primary = hex_to_rgb(palette["primary"])
    secondary = hex_to_rgb(palette["secondary"])
    text_dark = hex_to_rgb(palette["text_on_light"])

    set_background(slide, hex_to_rgb("FFFFFF"))

    title_text = data.get("title", "")
    title_size = fit_title_size(title_text, base_size=32, soft_limit=18, min_size=24)
    add_textbox(slide, Inches(0.8), Inches(0.55), Inches(11.5), Inches(1.0),
                title_text, title_size, text_dark, bold=True, font="Cambria")

    bullets = data.get("bullets", []) or []
    n = max(len(bullets), 1)
    margin = Inches(0.8)
    gap = Inches(0.3)
    col_w = (SLIDE_W - margin * 2 - gap * (n - 1)) / n if n > 0 else SLIDE_W - margin * 2
    circle_d = Inches(0.7)

    for i, b in enumerate(bullets):
        x = margin + (col_w + gap) * i
        circle_x = x + (col_w - circle_d) / 2
        add_circle_with_number(slide, circle_x, Inches(2.0), circle_d, i + 1, primary,
                                hex_to_rgb("FFFFFF"), font_size=22)
        add_textbox(slide, x, Inches(3.0), col_w, Inches(2.3), b, 17, text_dark,
                    font="Calibri", line_spacing=1.2, anchor=MSO_ANCHOR.TOP)

    # 画像専用エリア（本文テキストとは重ならない下段に固定）
    img_box = (Inches(0.8), Inches(5.55), Inches(2.6), Inches(1.65))
    add_picture_safe(slide, data.get("image_name"), *img_box)
    add_page_number(slide, idx, total, text_dark)
    return slide

# ============================================================
# レイアウト：統計／引用 強調スライド
# ============================================================
def layout_stat(prs, data, palette, idx, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    primary = hex_to_rgb(palette["primary"])
    secondary = hex_to_rgb(palette["secondary"])

    set_background(slide, hex_to_rgb(palette["text_on_light"]))
    add_circle(slide, Inches(-1.8), Inches(-1.8), Inches(4.0), primary)

    stat = data.get("stat", data.get("title", ""))
    label = data.get("bullets", [""])[0] if data.get("bullets") else ""
    stat_size = fit_title_size(stat, base_size=64, soft_limit=8, min_size=36, step=6)

    add_textbox(slide, Inches(1.0), Inches(2.2), Inches(11.3), Inches(2.0),
                stat, stat_size, secondary, bold=True, align=PP_ALIGN.CENTER, font="Cambria")
    add_textbox(slide, Inches(1.5), Inches(4.7), Inches(10.3), Inches(1.0),
                label, 18, hex_to_rgb("FFFFFF"), align=PP_ALIGN.CENTER, italic=True, font="Calibri")
    add_page_number(slide, idx, total, hex_to_rgb("FFFFFF"))
    return slide

# ============================================================
# レイアウト：まとめ／クロージング
# ============================================================
def layout_summary(prs, data, palette, idx, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    primary = hex_to_rgb(palette["primary"])
    secondary = hex_to_rgb(palette["secondary"])
    accent = hex_to_rgb(palette["accent"])

    set_background(slide, primary)
    add_circle(slide, SLIDE_W - Inches(2.5), SLIDE_H - Inches(2.5), Inches(4.0), secondary)

    title_text = data.get("title", "まとめ")
    title_size = fit_title_size(title_text, base_size=34, soft_limit=16, min_size=26)
    add_textbox(slide, Inches(1.0), Inches(0.9), Inches(10.5), Inches(1.0),
                title_text, title_size, accent, bold=True, font="Cambria")

    bullets = data.get("bullets", []) or []
    y = Inches(2.2)
    for b in bullets:
        add_rect(slide, Inches(1.0), y, Inches(0.12), Inches(0.55), secondary)
        add_textbox(slide, Inches(1.35), y, Inches(10.0), Inches(0.9), b, 19, accent, font="Calibri", line_spacing=1.15)
        y += Inches(1.0)

    add_page_number(slide, idx, total, accent)
    return slide

LAYOUTS = {
    "title": layout_title,
    "chapter": layout_chapter,
    "content_split": layout_content_split,
    "content_icons": layout_content_icons,
    "stat": layout_stat,
    "summary": layout_summary,
}

# ============================================================
# PPTX生成本体
# ============================================================
def create_pptx(payload):
    palette = get_palette(payload.get("palette", "midnight_executive"))
    slide_data_list = payload.get("slides", [])
    total = len(slide_data_list)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    for i, data in enumerate(slide_data_list, start=1):
        layout_key = data.get("layout", "content_split")
        layout_fn = LAYOUTS.get(layout_key, layout_content_split)
        layout_fn(prs, data, palette, i, total)

    ppt_stream = io.BytesIO()
    prs.save(ppt_stream)
    ppt_stream.seek(0)
    return ppt_stream

# ============================================================
# Streamlit UI
# ============================================================
st.title("🎨 Slide Generator AI PRO")
st.caption("テーマや資料を入れるだけで、配色・レイアウトにこだわった「映える」パワポを自動生成します。")

with st.sidebar:
    st.header("設定")
    num_slides = st.slider("スライド枚数（表紙・まとめ含む）", min_value=4, max_value=16, value=8)

uploaded_file = st.file_uploader("資料をアップロード（PDF / Word）任意", type=["pdf", "docx"])
theme = st.text_area("テーマや発表したい内容を入力してください：", height=140,
                      placeholder="例：PythonによるWebスクレイピングの基礎と注意点")

if st.button("スライドを生成する", type="primary"):
    context_text = ""
    if uploaded_file:
        with st.spinner("ファイルを解析中..."):
            context_text = extract_text(uploaded_file)

    source = (context_text + "\n" + theme).strip() if context_text else theme.strip()

    if source:
        with st.spinner("AIが構成とデザインを考えています…"):
            image_keywords = (
                "school_entrance, job_hunting, graph_bar_up, faces_four, study_man, study_man_smile, "
                "think_man, think_man_question, team_three, trash_bin, nausea_man, family_car, "
                "step_up_man, step_up_night, family_three, mouse_trap, cloud_upload_download, "
                "travel_bell, oni_two, oni_items, sakura_tree, graduation_cap, graph_cylinder, "
                "arrow_loop, sns_icons, presentation_man, earth_eco, smart_city, eco_factory, "
                "teacher_woman, dev_icons, network_cloud, meeting_two, meeting_room, trouble_man, "
                "graph_down_man, beginner_woman, umbrella_blue, umbrella_rain, farmers_two, "
                "river_flow, tanabata, susuki, rabbits_mochi, moon_rabbit, tsukimi_dango, "
                "ice_cubes, calendar_pages, calendar_clock, batsu_man, sincerity_man, sincerity_woman, "
                "pc_desk_man, pc_lap_man, think_couple, bubble_woman, think_cloud_woman, think_cloud_man, "
                "chat_two, chat_whisper, idea_man, magnifier_woman, home_energy, buildings, "
                "books_stack, smartphone_check, earth_simple, pc_desktop, bed_elderly, pc_desk_woman, calc_assets"
            )

            prompt = f"""
あなたはプロのプレゼンテーションデザイナーです。以下の【内容】を元に、合計{num_slides}枚のスライド構成データを
1つのJSONオブジェクトとして出力してください。余計な説明やMarkdownの```json囲みは一切不要、生のJSONのみを返してください。

【内容】
{source[:4000]}

【出力フォーマット】
{{
  "palette": "midnight_executive | forest_moss | coral_energy | ocean_gradient | charcoal_minimal | teal_trust | berry_cream | cherry_bold の中から内容に最も合うものを1つ",
  "slides": [
    {{
      "layout": "title",
      "title": "プレゼン全体のタイトル",
      "bullets": ["サブタイトルや発表者情報など1点"]
    }},
    {{
      "layout": "chapter",
      "chapter_no": "01",
      "title": "章のタイトル"
    }},
    {{
      "layout": "content_split または content_icons",
      "title": "スライドタイトル",
      "bullets": ["要点1", "要点2", "要点3"],
      "image_name": "上記キーワード一覧から1つ、なければ none"
    }},
    {{
      "layout": "stat",
      "stat": "強調したい数字やキーフレーズ（例：'95%' や '3倍'）",
      "bullets": ["その数字の説明1点"]
    }},
    {{
      "layout": "summary",
      "title": "まとめ",
      "bullets": ["まとめポイント1", "まとめポイント2", "まとめポイント3"]
    }}
  ]
}}

【構成ルール】
- 1枚目は必ず layout: "title"。
- 章が複数ある場合は、各章の最初に layout: "chapter" を入れる（chapter_noは"01","02"...）。
- 本文は layout: "content_split" と "content_icons" を交互に使い、同じレイアウトを連続させない。
- 内容の中に強調すべき数字・統計・キーフレーズがあれば layout: "stat" を1〜2枚挿入する。
- 最後のスライドは必ず layout: "summary"。
- 合計スライド数は{num_slides}枚に近づけること。

【文章の質に関する厳格なルール】
- title は12〜20文字程度。短すぎる単語だけや、長すぎる一文は禁止。
- 本文スライド（content_split / content_icons）の bullets は1スライドにつき3つ、各20〜40文字程度の「具体的な事実・固有名詞・数字を含む完全な文」にすること。
  - NG例（抽象的すぎる）：「魅力がある」「いろいろある」「素晴らしい体験」
  - OK例（具体的）：「銀山温泉に代表される、大正ロマンの街並みが残る名湯」
- 【内容】に書かれている情報から逸脱しない。情報が不足する場合は一般的な事実で補ってもよいが、無関係な内容や事実と異なる内容は書かないこと。
- stat の値は8文字以内の短いフレーズ（例："95%"、"年間300万人"）。bulletsにその数字の説明を1点、30文字程度で書く。
- summary の bullets は3〜4つ、各20〜30文字程度で、スライド全体の要点を要約すること（タイトルの単純な繰り返しは禁止）。

■ image_name に使えるキーワード一覧（内容に合うものを厳選、不要なら "none"）:
{image_keywords}
"""
            try:
                json_res = generate_with_retry(prompt)
                payload = json.loads(json_res)

                ppt_file = create_pptx(payload)

                st.success("🎉 スライドの生成に成功しました！下のボタンからダウンロードしてください。")
                st.download_button(
                    label="📥 パワーポイントファイルをダウンロード",
                    data=ppt_file,
                    file_name="generated_slides.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                )
            except Exception as e:
                st.error("エラーが発生しました。もう一度お試しください。")
                st.write(e)
    else:
        st.warning("テーマを入力するか、資料をアップロードしてください。")